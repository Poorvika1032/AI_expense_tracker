"""Flask application for the AI travel expense planning project."""
from datetime import date, datetime, timedelta, timezone
import csv
import calendar as calendar_module
from contextlib import contextmanager
import io
import json
import logging
import math
import os
import secrets
import hmac
import re
from pathlib import Path
import pickle
import sqlite3
import statistics
import uuid
from urllib.request import Request,urlopen
from PIL import Image, ImageOps

import pandas as pd
from flask import Flask, Response, abort, jsonify, redirect, render_template, request, session, url_for
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "database.db"
MODEL_PATH = BASE_DIR / "travel_expense_model.pkl"
DATA_PATH = BASE_DIR / "dataset" / "travel_expense.csv"

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
    MAX_CONTENT_LENGTH=5 * 1024 * 1024,
)
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


@contextmanager
def connect_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def create_database():
    """Create new tables and add columns safely to existing project databases."""
    with connect_db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL, password TEXT NOT NULL)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS travel_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER,
            source TEXT, destination TEXT, travelers INTEGER, days INTEGER,
            transportation TEXT, transportation_cost REAL, accommodation_cost REAL,
            food_cost REAL, other_expenses REAL, predicted_expense REAL, budget REAL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id))""")
        _add_column(conn, "users", "role", "TEXT NOT NULL DEFAULT 'USER'")
        _add_column(conn, "users", "is_active", "INTEGER NOT NULL DEFAULT 1")
        _add_column(conn, "users", "created_at", "TEXT")
        _add_column(conn, "travel_history", "created_at", "TEXT")
        _add_column(conn, "travel_history", "start_date", "TEXT")
        _add_column(conn, "travel_history", "end_date", "TEXT")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_travel_user ON travel_history(user_id, id DESC)")
        conn.execute("""CREATE TABLE IF NOT EXISTS expenses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trip_id INTEGER NOT NULL, user_id INTEGER NOT NULL,
            category TEXT NOT NULL, description TEXT DEFAULT '', amount REAL NOT NULL CHECK(amount >= 0),
            expense_date TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(trip_id) REFERENCES travel_history(id) ON DELETE CASCADE,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE)""")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_expenses_user_trip ON expenses(user_id, trip_id)")
        _add_column(conn, "users", "preferred_transport", "TEXT")
        _add_column(conn, "users", "default_budget", "REAL")
        _add_column(conn, "users", "currency", "TEXT NOT NULL DEFAULT 'INR'")
        _add_column(conn, "users", "theme", "TEXT NOT NULL DEFAULT 'light'")
        _add_column(conn, "users", "profile_image", "TEXT")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_expenses_date ON expenses(user_id, expense_date)")
        conn.execute("""CREATE TABLE IF NOT EXISTS wishlist (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
            destination TEXT NOT NULL, estimated_cost REAL, target_date TEXT, notes TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
            UNIQUE(user_id,destination))""")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_wishlist_user ON wishlist(user_id,id DESC)")
        conn.execute("""CREATE TABLE IF NOT EXISTS favorite_destinations (
            user_id INTEGER NOT NULL,destination TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(user_id,destination),
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS user_savings (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, trip_id INTEGER,
            original_cost REAL NOT NULL, optimized_cost REAL NOT NULL, saved_amount REAL NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY(trip_id) REFERENCES travel_history(id) ON DELETE SET NULL)""")
        _add_column(conn,"user_savings","actual_cost","REAL")
        _add_column(conn,"user_savings","actual_savings","REAL")
        conn.execute("""CREATE TABLE IF NOT EXISTS notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
            title TEXT NOT NULL, message TEXT NOT NULL, notification_type TEXT DEFAULT 'info',
            is_read INTEGER NOT NULL DEFAULT 0, created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS datasets (
            id INTEGER PRIMARY KEY AUTOINCREMENT, stored_name TEXT UNIQUE NOT NULL,
            original_name TEXT NOT NULL, row_count INTEGER NOT NULL, uploaded_by INTEGER,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(uploaded_by) REFERENCES users(id))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS exchange_rates (
            base TEXT NOT NULL, quote TEXT NOT NULL, rate REAL NOT NULL,
            rate_date TEXT NOT NULL, fetched_at TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(base,quote,rate_date))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS activity_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER,
            action TEXT NOT NULL, details TEXT DEFAULT '', created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS model_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, version TEXT, model_path TEXT,
            training_date TEXT, dataset_size INTEGER, mae REAL, mse REAL, r2_score REAL)""")
        _add_column(conn, "model_versions", "rmse", "REAL")
        _add_column(conn, "model_versions", "notes", "TEXT")
        _add_column(conn, "model_versions", "is_active", "INTEGER NOT NULL DEFAULT 0")


def _add_column(conn, table, column, definition):
    columns = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}
    if column not in columns:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def refresh_actual_savings(conn,user_id,trip_id):
    actual=conn.execute("SELECT SUM(amount) FROM expenses WHERE user_id=? AND trip_id=?",(user_id,trip_id)).fetchone()[0]
    conn.execute("""UPDATE user_savings SET actual_cost=?,
        actual_savings=CASE WHEN ? IS NULL THEN NULL ELSE original_cost-? END
        WHERE user_id=? AND trip_id=?""",(actual,actual,actual,user_id,trip_id))


create_database()


@app.cli.command("create-admin")
def create_admin():
    """Create or promote an administrator without exposing credentials in logs."""
    import getpass
    name = input("Admin name: ").strip()
    email = input("Admin email: ").strip().lower()
    password = getpass.getpass("Admin password (8+ characters): ")
    if not name or len(name) > 100 or "@" not in email or len(password) < 8:
        raise SystemExit("Enter a valid name/email and a password of at least 8 characters.")
    with connect_db() as conn:
        user = conn.execute("SELECT id FROM users WHERE lower(email)=?", (email,)).fetchone()
        if user:
            conn.execute("UPDATE users SET role='ADMIN',is_active=1 WHERE id=?", (user["id"],))
        else:
            conn.execute("INSERT INTO users(name,email,password,role,is_active) VALUES(?,?,?,?,1)",
                         (name, email, generate_password_hash(password), "ADMIN"))
    print("Administrator account is ready.")


@app.context_processor
def inject_csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"]}


@app.before_request
def protect_state_changes():
    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        expected = session.get("csrf_token", "")
        payload=request.get_json(silent=True)
        if not isinstance(payload,dict): payload={}
        supplied = (request.headers.get("X-CSRF-Token") or request.form.get("csrf_token", "")
                    or payload.get("csrf_token", ""))
        if not expected or not hmac.compare_digest(expected, supplied):
            abort(400, "The form expired. Reload the page and try again.")


def load_model():
    try:
        artifact_path = MODEL_PATH
        with connect_db() as conn:
            active = conn.execute("SELECT model_path FROM model_versions WHERE is_active=1 LIMIT 1").fetchone()
        if active and active["model_path"]:
            candidate = (BASE_DIR / active["model_path"]).resolve()
            models_dir = (BASE_DIR / "models").resolve()
            if candidate.parent == models_dir and candidate.is_file():
                artifact_path = candidate
        with artifact_path.open("rb") as file:
            saved = pickle.load(file)
        # New artifacts carry the complete sklearn preprocessing pipeline.
        if isinstance(saved, dict) and "pipeline" in saved:
            metrics = dict(saved.get("metrics", {}))
            metrics.setdefault("trained_at", saved.get("trained_at"))
            return saved["pipeline"], metrics
        # Older project artifacts used category integers and cannot safely encode unseen cities.
        # Fall back to training the upgraded pipeline when dependencies/data are available.
        from train_model import train
        train()
        with artifact_path.open("rb") as file:
            saved = pickle.load(file)
        metrics = dict(saved.get("metrics", {}))
        metrics.setdefault("trained_at", saved.get("trained_at"))
        return saved["pipeline"], metrics
    except Exception:
        logger.exception("Could not load the travel model")
        return None, {}


model, model_metrics = load_model()


def authenticated():
    return session.get("user_id") is not None


def owned_trip(trip_id, user_id):
    with connect_db() as conn:
        return conn.execute("SELECT id FROM travel_history WHERE id=? AND user_id=?", (trip_id, user_id)).fetchone()


@app.route("/")
def home():
    return redirect(url_for("dashboard")) if authenticated() else redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html")
    email, password = request.form.get("email", "").strip().lower(), request.form.get("password", "")
    with connect_db() as conn:
        user = conn.execute("SELECT * FROM users WHERE lower(email)=? AND is_active=1", (email,)).fetchone()
    if user and check_password_hash(user["password"], password):
        with connect_db() as conn:
            conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                         (user["id"],"login","Successful login"))
        session.clear()
        session["user_id"], session["user"] = user["id"], user["name"]
        session["role"] = user["role"]
        session["theme"] = user["theme"] or "light"
        logger.info("User login: id=%s", user["id"])
        return redirect(url_for("dashboard"))
    return render_template("login.html", error="Invalid email or password."), 401


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "GET":
        return render_template("register.html")
    name = request.form.get("name", "").strip()
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")
    if not name or len(name) > 100 or "@" not in email or len(password) < 8:
        return render_template("register.html", error="Enter a valid name and email; password must be at least 8 characters."), 400
    try:
        with connect_db() as conn:
            conn.execute("INSERT INTO users(name,email,password) VALUES(?,?,?)",
                         (name, email, generate_password_hash(password)))
    except sqlite3.IntegrityError:
        return render_template("register.html", error="That email is already registered."), 409
    with connect_db() as conn:
        created=conn.execute("SELECT id FROM users WHERE lower(email)=?",(email,)).fetchone()
        conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                     (created["id"],"registration","Account created"))
    return redirect(url_for("login"))


@app.route("/dashboard")
def dashboard():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        account=conn.execute("SELECT id FROM users WHERE id=? AND is_active=1",(session["user_id"],)).fetchone()
        if not account:
            session.clear()
            return redirect(url_for("login"))
        trips = conn.execute("SELECT * FROM travel_history WHERE user_id=? ORDER BY id DESC LIMIT 5",
                             (session["user_id"],)).fetchall()
        stats = conn.execute("""SELECT COUNT(*) trips, COALESCE(SUM(predicted_expense),0) total,
            COALESCE(AVG(predicted_expense),0) average,
            COALESCE(MAX(predicted_expense),0) highest,
            COALESCE(MIN(predicted_expense),0) lowest
            FROM travel_history WHERE user_id=?""",
                             (session["user_id"],)).fetchone()
        actual = conn.execute("SELECT COALESCE(SUM(amount),0) FROM expenses WHERE user_id=?",
                              (session["user_id"],)).fetchone()[0]
        savings = conn.execute("SELECT COALESCE(SUM(saved_amount),0) FROM user_savings WHERE user_id=?",
                               (session["user_id"],)).fetchone()[0]
        actual_savings = conn.execute("SELECT COALESCE(SUM(actual_savings),0) FROM user_savings WHERE user_id=?",
                                      (session["user_id"],)).fetchone()[0]
        favorites = conn.execute("""SELECT destination,COUNT(*) visits FROM travel_history
            WHERE user_id=? GROUP BY destination ORDER BY visits DESC LIMIT 1""",
                                 (session["user_id"],)).fetchone()
        transport = conn.execute("""SELECT transportation,COUNT(*) uses FROM travel_history
            WHERE user_id=? GROUP BY transportation ORDER BY uses DESC LIMIT 1""",
                                 (session["user_id"],)).fetchone()
        upcoming=conn.execute("""SELECT id,source,destination,start_date FROM travel_history
            WHERE user_id=? AND start_date BETWEEN ? AND ? ORDER BY start_date LIMIT 5""",
            (session["user_id"],date.today().isoformat(),(date.today()+timedelta(days=7)).isoformat())).fetchall()
        for trip in upcoming:
            message=f"Trip {trip['id']} to {trip['destination']} starts on {trip['start_date']}."
            try:
                conn.execute("""INSERT INTO notifications(user_id,title,message,notification_type)
                    SELECT ?,?,?,? WHERE NOT EXISTS(SELECT 1 FROM notifications WHERE user_id=? AND title=? AND message=?)""",
                    (session["user_id"],"Trip approaching",message,"trip",session["user_id"],"Trip approaching",message))
            except sqlite3.Error:
                logger.exception("Could not create upcoming-trip notification for user_id=%s",session["user_id"])
    return render_template("dashboard.html", recent_trips=trips, trip_count=stats["trips"],
                           total_spend=round(stats["total"], 2), average_spend=round(stats["average"], 2),
                           actual_spend=round(actual, 2), total_savings=round(savings,2), actual_savings=round(actual_savings,2), highest=stats["highest"], lowest=stats["lowest"],
                           upcoming_trips=upcoming,
                           favorite=favorites["destination"] if favorites else "-",
                           common_transport=transport["transportation"] if transport else "-",
                           model_available=model is not None)


def city_name(value):
    value = " ".join(value.split()).title()
    return "Bangalore" if value == "Bengaluru" else value


def estimate_trip(source, destination, travelers, days, transports=("Bus", "Train", "Car"), predictor=None):
    """Return model-based component and total estimates for each requested mode."""
    predictor = predictor or model
    if predictor is None:
        raise RuntimeError("The prediction model is unavailable.")
    predictions = {}
    for transport in transports:
        frame = pd.DataFrame([{"source": source, "destination": destination,
                               "transportation": transport, "travelers": travelers, "days": days}])
        components = [max(0.0, round(float(value), 2)) for value in predictor.predict(frame)[0]]
        predictions[transport] = dict(zip(
            ["transport_cost", "accommodation_cost", "food_cost", "other_expenses"], components))
        predictions[transport]["total_cost"] = round(sum(components), 2)
    return predictions


def require_admin():
    if not authenticated():
        return redirect(url_for("login"))
    if session.get("role") != "ADMIN":
        abort(403)
    return None


@app.route("/predict", methods=["POST"])
def predict():
    if not authenticated():
        return redirect(url_for("login"))
    try:
        source = city_name(request.form.get("source", ""))
        destination = city_name(request.form.get("destination", ""))
        travelers = int(request.form.get("travelers", ""))
        days = int(request.form.get("days", ""))
        start_raw = request.form.get("start_date", "").strip()
        end_raw = request.form.get("end_date", "").strip()
        start_date = date.fromisoformat(start_raw) if start_raw else None
        end_date = date.fromisoformat(end_raw) if end_raw else None
        if bool(start_date) != bool(end_date) or (start_date and end_date < start_date):
            raise ValueError("Enter a valid start and end date.")
        if start_date:
            days = (end_date - start_date).days + 1
        transport_choice = request.form.get("transportation", "")
        budget = float(request.form.get("budget", ""))
        if not source or not destination or len(source) > 100 or len(destination) > 100 or source == destination:
            raise ValueError("Enter two different places.")
        if not (1 <= travelers <= 100) or not (1 <= days <= 365):
            raise ValueError("Travelers must be 1-100 and trip duration 1-365 days.")
        if transport_choice not in {"Bus", "Train", "Car"} or not math.isfinite(budget) or not (0 < budget <= 1_000_000_000):
            raise ValueError("Choose a supported transport and enter a positive budget.")
    except (ValueError, TypeError):
        return render_template("dashboard.html", error="Please check the trip details and enter valid values."), 400
    if model is None:
        return render_template("dashboard.html", error="The prediction model is unavailable. Train it with train_model.py and try again."), 503

    predictions = estimate_trip(source, destination, travelers, days)
    within_budget = {name: values for name, values in predictions.items() if values["total_cost"] <= budget}
    best_transport = (transport_choice if transport_choice in within_budget else
                      min(within_budget, key=lambda item: predictions[item]["total_cost"]) if within_budget else transport_choice)
    best = predictions[best_transport]
    recommendation = (f"Random Forest estimates {best_transport} at â‚¹{best['total_cost']:.2f}; "
                      + (f"â‚¹{budget-best['total_cost']:.2f} remains in budget." if best_transport in within_budget
                         else "all transport estimates exceed your budget."))
    with connect_db() as conn:
        conn.execute("""INSERT INTO travel_history
            (user_id,source,destination,travelers,days,transportation,transportation_cost,
             accommodation_cost,food_cost,other_expenses,predicted_expense,budget,start_date,end_date)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (session["user_id"], source, destination, travelers, days, best_transport,
             best["transport_cost"], best["accommodation_cost"], best["food_cost"],
             best["other_expenses"], best["total_cost"], budget,
             start_date.isoformat() if start_date else None,end_date.isoformat() if end_date else None))
        conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                     (session["user_id"], "prediction", f"{source} to {destination}"))
        if best["total_cost"] > budget:
            conn.execute("INSERT INTO notifications(user_id,title,message,notification_type) VALUES(?,?,?,?)",
                         (session["user_id"], "Plan exceeds budget",
                          f"Estimated cost {best['total_cost']:.2f} exceeds budget {budget:.2f}.", "budget"))
    return render_template("dashboard.html", prediction=best["total_cost"],
        transportation_cost=best["transport_cost"], accommodation_cost=best["accommodation_cost"],
        food_cost=best["food_cost"], other_expenses=best["other_expenses"], budget=budget,
        recommendation=recommendation, predictions=predictions, trip_count=None)


@app.route("/history")
def history():
    if not authenticated():
        return redirect(url_for("login"))
    try: page=max(1,min(10000,int(request.args.get("page",1) or 1)))
    except ValueError: page=1
    per_page=10
    clauses=["user_id=?"]; params=[session["user_id"]]
    destination=request.args.get("destination","").strip()[:100]
    transport=request.args.get("transportation","").strip()
    status=request.args.get("budget_status","").strip()
    search=request.args.get("q","").strip()[:100]
    date_from=request.args.get("date_from","").strip()
    date_to=request.args.get("date_to","").strip()
    expense_min=request.args.get("expense_min","").strip()
    expense_max=request.args.get("expense_max","").strip()
    sort=request.args.get("sort","newest")
    if destination: clauses.append("destination LIKE ?");params.append(f"%{destination}%")
    if search: clauses.append("(source LIKE ? OR destination LIKE ?)");params.extend([f"%{search}%",f"%{search}%"])
    if transport in {"Bus","Train","Car"}: clauses.append("transportation=?");params.append(transport)
    if status=="within": clauses.append("predicted_expense<=budget")
    elif status=="over": clauses.append("predicted_expense>budget")
    if date_from:
        try: date.fromisoformat(date_from);clauses.append("date(created_at)>=date(?)");params.append(date_from)
        except ValueError: date_from=""
    if date_to:
        try: date.fromisoformat(date_to);clauses.append("date(created_at)<=date(?)");params.append(date_to)
        except ValueError: date_to=""
    for raw,operator in ((expense_min,">="),(expense_max,"<=")):
        if raw:
            try:
                amount=float(raw)
                if not math.isfinite(amount) or amount<0: raise ValueError
                clauses.append(f"predicted_expense{operator}?");params.append(amount)
            except ValueError: pass
    order={"newest":"id DESC","oldest":"id ASC","cost_high":"predicted_expense DESC","cost_low":"predicted_expense ASC","destination":"destination COLLATE NOCASE ASC"}.get(sort,"id DESC")
    where=" AND ".join(clauses)
    with connect_db() as conn:
        total=conn.execute(f"SELECT COUNT(*) FROM travel_history WHERE {where}",params).fetchone()[0]
        rows=conn.execute(f"""SELECT source,destination,travelers,days,transportation,transportation_cost,
            accommodation_cost,food_cost,other_expenses,predicted_expense,budget,id
            FROM travel_history WHERE {where} ORDER BY {order} LIMIT ? OFFSET ?""",params+[per_page,(page-1)*per_page]).fetchall()
    return render_template("history.html",history=rows,page=page,pages=max(1,(total+per_page-1)//per_page),total=total)


@app.route("/insights")
def insights():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        row=conn.execute("""SELECT AVG(predicted_expense) average,AVG(days) days,
            (SELECT transportation FROM travel_history WHERE user_id=? GROUP BY transportation ORDER BY COUNT(*) DESC LIMIT 1) transport
            FROM travel_history WHERE user_id=?""",(session["user_id"],session["user_id"])).fetchone()
    average=row["average"] or 0
    message=("No travel history is available yet." if average==0 else
             f"Your recorded trips average {row['days']:.1f} days. You most often select {row['transport'] or 'no transport'}; compare alternatives on each plan.")
    return render_template("insights.html", average_expense=round(average, 2), message=message,
                           average_days=round(row["days"] or 0,1),preferred_transport=row["transport"] or "Not enough data")


@app.route("/reports/export.csv")
def export_csv():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        rows = conn.execute("""SELECT t.source,t.destination,t.travelers,t.days,t.transportation,t.predicted_expense,t.budget,t.created_at,
            (SELECT SUM(e.amount) FROM expenses e WHERE e.trip_id=t.id AND e.user_id=t.user_id) actual_expense
            FROM travel_history t WHERE t.user_id=? ORDER BY t.id DESC""",
                            (session["user_id"],)).fetchall()
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(["Source", "Destination", "Travelers", "Days", "Transport", "Predicted expense", "Budget", "Created", "Actual expense"])
    def safe_cell(value):
        if isinstance(value,str) and value[:1] in {"=","+","-","@"}:
            return "'"+value
        return value
    writer.writerows([[safe_cell(value) for value in row] for row in rows])
    return Response(stream.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=travel-report.csv"})


@app.route("/reports/export.pdf")
def export_pdf():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        rows=conn.execute("""SELECT t.source,t.destination,t.days,t.travelers,t.predicted_expense,
            (SELECT SUM(e.amount) FROM expenses e WHERE e.trip_id=t.id AND e.user_id=t.user_id) actual
            FROM travel_history t WHERE t.user_id=? ORDER BY t.id DESC LIMIT 35""",(session["user_id"],)).fetchall()
    lines=["AI TRAVEL EXPENSE REPORT",f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}","Amounts shown in INR","", "DESTINATION | DAYS | PEOPLE | PREDICTED | ACTUAL"]
    for row in rows:
        actual=f"{row['actual']:.2f}" if row["actual"] is not None else "Not recorded"
        lines.append(f"{row['source']} to {row['destination']} | {row['days']} | {row['travelers']} | {row['predicted_expense']:.2f} | {actual}")
    if not rows: lines.append("No trips recorded.")
    commands=["BT /F1 11 Tf 48 760 Td"]
    for index,line in enumerate(lines[:40]):
        safe=line.encode("ascii","replace").decode("ascii").replace("\\","\\\\").replace("(","\\(").replace(")","\\)")[:115]
        commands.append(f"({safe}) Tj")
        if index<len(lines[:40])-1: commands.append("0 -17 Td")
    stream=" ".join(commands)+" ET"
    objects=[b"<< /Type /Catalog /Pages 2 0 R >>",b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",f"<< /Length {len(stream.encode('ascii'))} >>\nstream\n{stream}\nendstream".encode("ascii")]
    pdf=bytearray(b"%PDF-1.4\n")
    offsets=[0]
    for num,obj in enumerate(objects,1):
        offsets.append(len(pdf));pdf.extend(f"{num} 0 obj\n".encode());pdf.extend(obj);pdf.extend(b"\nendobj\n")
    xref=len(pdf);pdf.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return Response(bytes(pdf),mimetype="application/pdf",headers={"Content-Disposition":"attachment; filename=travel-report.pdf"})


@app.route("/logout", methods=["POST"])
def logout():
    user_id=session.get("user_id")
    if user_id:
        with connect_db() as conn:
            conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",(user_id,"logout","Session ended"))
    session.clear()
    return redirect(url_for("login"))


@app.errorhandler(404)
def not_found(_error):
    return render_template("error.html",code=404,message="We could not find that page."),404


@app.errorhandler(403)
def forbidden(_error):
    return render_template("error.html",code=403,message="You do not have access to this page."),403


@app.errorhandler(400)
def bad_request(error):
    if request.path.startswith("/api/"):
        return jsonify({"success":False,"error":getattr(error,"description","Invalid request.")}),400
    return render_template("error.html",code=400,message=getattr(error,"description","Please check your request.")),400


@app.errorhandler(500)
def server_error(error):
    original = getattr(error, "original_exception", None)
    location = f"{request.method} {request.path} user_id={session.get('user_id', 'anonymous')}"
    if original is not None:
        logger.error("Unhandled server error on %s", location, exc_info=(type(original), original, original.__traceback__))
    else:
        logger.exception("Unhandled HTTP 500 error on %s: %s", location, error)
    debug_mode = app.debug or os.environ.get("FLASK_DEBUG") == "1"
    if debug_mode:
        detail = str(original) if original is not None else str(error)
        return render_template("error.html", code=500, message=f"Server error: {detail}"), 500
    return render_template("error.html", code=500, message="Something went wrong. Please try again."), 500


@app.errorhandler(413)
def upload_too_large(_error):
    return render_template("error.html",code=413,message="The uploaded file exceeds the 5 MB limit."),413


@app.route("/trips/<int:trip_id>")
def trip_detail(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        trip = conn.execute("SELECT * FROM travel_history WHERE id=? AND user_id=?",
                             (trip_id, session["user_id"])).fetchone()
        if not trip:
            return "Trip not found", 404
        expenses = conn.execute("SELECT * FROM expenses WHERE trip_id=? AND user_id=? ORDER BY expense_date DESC, id DESC",
                                (trip_id, session["user_id"])).fetchall()
        favorite=conn.execute("SELECT 1 FROM favorite_destinations WHERE user_id=? AND destination=?",
                              (session["user_id"],trip["destination"])).fetchone() is not None
        daily=conn.execute("SELECT expense_date,SUM(amount) amount FROM expenses WHERE trip_id=? AND user_id=? GROUP BY expense_date ORDER BY expense_date",
                           (trip_id,session["user_id"])).fetchall()
    actual = sum(row["amount"] for row in expenses)
    return render_template("trip_detail.html", trip=trip, expenses=expenses,
                           actual=round(actual, 2), difference=round(actual-trip["predicted_expense"], 2),favorite=favorite,
                           daily_expenses=daily,daily_target=round((trip["predicted_expense"] or 0)/max(1,trip["days"] or 1),2))


@app.route("/trips/<int:trip_id>/expenses", methods=["POST"])
def add_expense(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    category = request.form.get("category", "").strip()
    description = request.form.get("description", "").strip()[:250]
    try:
        amount = float(request.form.get("amount", ""))
        expense_date = date.fromisoformat(request.form.get("expense_date", "" )).isoformat()
        if category not in {"Transportation", "Accommodation", "Food", "Activities", "Shopping", "Other", "Emergency"} or not (0 < amount <= 1_000_000_000):
            raise ValueError
    except (ValueError, TypeError):
        return "Enter a valid expense category, positive amount, and date.", 400
    with connect_db() as conn:
        trip = conn.execute("SELECT id FROM travel_history WHERE id=? AND user_id=?",
                            (trip_id, session["user_id"])).fetchone()
        if not trip:
            return "Trip not found", 404
        previous=[float(row[0]) for row in conn.execute("SELECT amount FROM expenses WHERE user_id=? ORDER BY id DESC LIMIT 100",
                                                       (session["user_id"],)).fetchall()]
        unusual=False
        if len(previous)>=5:
            center=statistics.median(previous)
            deviation=statistics.median([abs(value-center) for value in previous])
            threshold=center+3*deviation if deviation>0 else center*2.5
            unusual=amount>threshold
        conn.execute("INSERT INTO expenses(trip_id,user_id,category,description,amount,expense_date) VALUES(?,?,?,?,?,?)",
                     (trip_id, session["user_id"], category, description, amount, expense_date))
        refresh_actual_savings(conn,session["user_id"],trip_id)
        if unusual:
            conn.execute("INSERT INTO notifications(user_id,title,message,notification_type) VALUES(?,?,?,?)",
                (session["user_id"],"Unusual expense detected",
                 f"This expense of {amount:.2f} is significantly higher than your historical travel spending (median {center:.2f}).","anomaly"))
    return redirect(url_for("trip_detail", trip_id=trip_id))


@app.route("/expenses/<int:expense_id>/delete", methods=["POST"])
def delete_expense(expense_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        row = conn.execute("SELECT trip_id FROM expenses WHERE id=? AND user_id=?",
                           (expense_id, session["user_id"])).fetchone()
        if not row:
            return "Expense not found", 404
        conn.execute("DELETE FROM expenses WHERE id=? AND user_id=?", (expense_id, session["user_id"]))
        refresh_actual_savings(conn,session["user_id"],row["trip_id"])
    return redirect(url_for("trip_detail", trip_id=row["trip_id"]))


@app.route("/analytics")
def analytics():
    if not authenticated():
        return redirect(url_for("login"))
    date_from=request.args.get("date_from","").strip();date_to=request.args.get("date_to","").strip()
    try:
        if date_from: date.fromisoformat(date_from)
        if date_to: date.fromisoformat(date_to)
        if date_from and date_to and date_from>date_to: raise ValueError
    except ValueError: date_from=date_to=""
    with connect_db() as conn:
        rows = conn.execute("""SELECT t.id,t.destination,t.predicted_expense,SUM(e.amount) actual
            FROM travel_history t LEFT JOIN expenses e ON e.trip_id=t.id AND e.user_id=t.user_id
            WHERE t.user_id=? AND (?='' OR date(t.created_at)>=date(?)) AND (?='' OR date(t.created_at)<=date(?))
            GROUP BY t.id ORDER BY t.id""", (session["user_id"],date_from,date_from,date_to,date_to)).fetchall()
    pairs = [(float(row["predicted_expense"] or 0), float(row["actual"])) for row in rows if row["actual"] is not None and row["actual"] > 0]
    errors = [abs(actual-predicted) for predicted, actual in pairs]
    mae = sum(errors)/len(errors) if errors else None
    rmse = (sum(error**2 for error in [actual-predicted for predicted,actual in pairs])/len(pairs))**0.5 if pairs else None
    return render_template("analytics.html", rows=rows, pair_count=len(pairs), mae=mae, rmse=rmse,
                           over=sum(actual > predicted for predicted,actual in pairs),
                           under=sum(actual < predicted for predicted,actual in pairs),date_from=date_from,date_to=date_to)


@app.route("/model-performance")
def model_performance():
    if not authenticated():
        return redirect(url_for("login"))
    metrics = model_metrics or {}
    with connect_db() as conn:
        active=conn.execute("SELECT version FROM model_versions WHERE is_active=1 LIMIT 1").fetchone()
    return render_template("model_performance.html", metrics=metrics,
                           rmse=metrics.get("rmse", (metrics.get("mse", 0) ** 0.5 if metrics.get("mse") is not None else None)),
                           trained_at=metrics.get("trained_at", "Stored in model artifact"),
                           version=active["version"] if active else "Bundled artifact")


@app.route("/model-insights")
def model_insights():
    if not authenticated():
        return redirect(url_for("login"))
    factors = []
    if model is not None:
        try:
            names = model.named_steps["preprocessing"].get_feature_names_out()
            importance = model.named_steps["regressor"].feature_importances_
            grouped = {}
            for name, score in zip(names, importance):
                # Aggregate encoded city/category columns back to their original feature.
                key = str(name).split("__", 1)[-1].split("_", 1)[0]
                grouped[key] = grouped.get(key, 0.0) + float(score)
            factors = sorted(grouped.items(), key=lambda item: item[1], reverse=True)[:10]
        except (AttributeError, KeyError, ValueError):
            logger.exception("Unable to calculate model feature importance")
    return render_template("model_insights.html", factors=factors)


@app.route("/admin")
def admin_dashboard():
    if not authenticated():
        return redirect(url_for("login"))
    if session.get("role") != "ADMIN":
        return "Forbidden", 403
    with connect_db() as conn:
        stats = conn.execute("SELECT (SELECT COUNT(*) FROM users) users,(SELECT COUNT(*) FROM travel_history) trips,(SELECT COUNT(*) FROM expenses) expenses").fetchone()
        destinations=conn.execute("SELECT destination,COUNT(*) visits FROM travel_history GROUP BY destination ORDER BY visits DESC LIMIT 5").fetchall()
        transports=conn.execute("SELECT transportation,COUNT(*) uses FROM travel_history GROUP BY transportation ORDER BY uses DESC LIMIT 5").fetchall()
        active=conn.execute("SELECT version,training_date,dataset_size,mae,rmse,r2_score FROM model_versions WHERE is_active=1 LIMIT 1").fetchone()
    return render_template("admin.html", stats=stats,destinations=destinations,transports=transports,active_model=active)

@app.route("/expenses/<int:expense_id>/edit", methods=["POST"])
def edit_expense(expense_id):
    if not authenticated():
        return redirect(url_for("login"))
    category = request.form.get("category", "").strip()
    description = request.form.get("description", "").strip()[:250]
    try:
        amount = float(request.form.get("amount", ""))
        expense_date = date.fromisoformat(request.form.get("expense_date", "")).isoformat()
        if category not in {"Transportation", "Accommodation", "Food", "Activities", "Shopping", "Other", "Emergency"} or not math.isfinite(amount) or not (0 < amount <= 1_000_000_000):
            raise ValueError
    except (ValueError, TypeError):
        return "Enter a valid expense category, positive amount, and date.", 400
    with connect_db() as conn:
        row = conn.execute("SELECT trip_id FROM expenses WHERE id=? AND user_id=?",
                           (expense_id, session["user_id"])).fetchone()
        if not row:
            return "Expense not found", 404
        conn.execute("UPDATE expenses SET category=?,description=?,amount=?,expense_date=? WHERE id=? AND user_id=?",
                     (category, description, amount, expense_date, expense_id, session["user_id"]))
        refresh_actual_savings(conn,session["user_id"],row["trip_id"])
    return redirect(url_for("trip_detail", trip_id=row["trip_id"]))
@app.route("/wishlist", methods=["GET", "POST"])
def wishlist():
    if not authenticated():
        return redirect(url_for("login"))
    error = None
    if request.method == "POST":
        destination = city_name(request.form.get("destination", ""))
        notes = request.form.get("notes", "").strip()[:300]
        target_date = request.form.get("target_date", "").strip() or None
        try:
            estimated_cost = float(request.form.get("estimated_cost", "")) if request.form.get("estimated_cost", "").strip() else None
            if not destination or len(destination) > 100 or (estimated_cost is not None and (not math.isfinite(estimated_cost) or estimated_cost < 0)):
                raise ValueError
            if target_date:
                target_date = date.fromisoformat(target_date).isoformat()
            with connect_db() as conn:
                conn.execute("INSERT INTO wishlist(user_id,destination,estimated_cost,target_date,notes) VALUES(?,?,?,?,?)",
                             (session["user_id"], destination, estimated_cost, target_date, notes))
        except sqlite3.IntegrityError:
            error = "That destination is already on your wishlist."
        except (ValueError, TypeError):
            error = "Enter a destination, valid target date, and non-negative estimated cost."
    with connect_db() as conn:
        rows = conn.execute("SELECT * FROM wishlist WHERE user_id=? ORDER BY id DESC", (session["user_id"],)).fetchall()
    return render_template("wishlist.html", items=rows, error=error)


@app.route("/wishlist/<int:item_id>/delete", methods=["POST"])
def delete_wishlist_item(item_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        conn.execute("DELETE FROM wishlist WHERE id=? AND user_id=?", (item_id, session["user_id"]))
    return redirect(url_for("wishlist"))


@app.route("/recommendations", methods=["GET", "POST"])
def destination_recommendations():
    if not authenticated():
        return redirect(url_for("login"))
    results = []
    error = None
    if request.method == "POST":
        try:
            source = city_name(request.form.get("source", ""))
            travelers = int(request.form.get("travelers", ""))
            days = int(request.form.get("days", ""))
            budget = float(request.form.get("budget", ""))
            preferred = request.form.get("transportation", "Train")
            if not source or len(source) > 100 or not 1 <= travelers <= 100 or not 1 <= days <= 365 or not math.isfinite(budget) or not 0 < budget <= 1_000_000_000 or preferred not in {"Bus", "Train", "Car"}:
                raise ValueError
            if model is None:
                raise RuntimeError("Prediction model unavailable.")
            available = sorted({city_name(value) for value in pd.read_csv(DATA_PATH)["destination"].dropna().astype(str) if city_name(value) != source})
            for destination in available:
                estimates = {}
                for transport in ["Bus", "Train", "Car"]:
                    features = pd.DataFrame([{"source": source, "destination": destination, "transportation": transport, "travelers": travelers, "days": days}])
                    estimate = max(0.0, float(model.predict(features)[0].sum()))
                    estimates[transport] = estimate
                chosen_transport = min(estimates, key=lambda name: (estimates[name], name != preferred))
                results.append({"destination": destination, "transport": chosen_transport,
                               "cost": round(estimates[chosen_transport], 2), "within_budget": estimates[chosen_transport] <= budget})
            results.sort(key=lambda item: (not item["within_budget"], item["cost"]))
        except (ValueError, TypeError, KeyError, OSError) as exc:
            logger.info("Destination recommendation input/data error: %s", type(exc).__name__)
            error = "Enter valid trip details; dataset destinations must be available to generate recommendations."
        except RuntimeError as exc:
            error = str(exc)
    return render_template("recommendations.html", results=results, error=error)

@app.route("/what-if", methods=["GET", "POST"])
def what_if():
    if not authenticated():
        return redirect(url_for("login"))
    results, error, base = [], None, None
    if request.method == "POST":
        try:
            source, destination = city_name(request.form.get("source", "")), city_name(request.form.get("destination", ""))
            travelers, days = int(request.form.get("travelers", "")), int(request.form.get("days", ""))
            budget = float(request.form.get("budget", ""))
            transport = request.form.get("transportation", "Train")
            if not source or not destination or source == destination or not 1 <= travelers <= 100 or not 1 <= days <= 365 or not math.isfinite(budget) or budget <= 0 or transport not in {"Bus", "Train", "Car"}:
                raise ValueError
            base = {"source": source, "destination": destination, "travelers": travelers, "days": days,
                    "transportation": transport, "budget": budget}
            scenarios = [("Current plan", days, travelers, transport),
                         ("One fewer day", max(1, days-1), travelers, transport),
                         ("Two fewer days", max(1, days-2), travelers, transport),
                         ("Bus", days, travelers, "Bus"), ("Train", days, travelers, "Train"),
                         ("Car", days, travelers, "Car"),
                         ("One fewer traveler", days, max(1, travelers-1), transport)]
            seen = set()
            for label, scenario_days, scenario_travelers, scenario_transport in scenarios:
                key = (scenario_days, scenario_travelers, scenario_transport)
                if key in seen:
                    continue
                seen.add(key)
                estimate = estimate_trip(source, destination, scenario_travelers, scenario_days, (scenario_transport,))[scenario_transport]["total_cost"]
                results.append({"label": label, "days": scenario_days, "travelers": scenario_travelers,
                                "transport": scenario_transport, "cost": estimate, "budget": budget})
        except (ValueError, TypeError, RuntimeError):
            error = "Enter valid trip details and ensure a model is available."
    return render_template("what_if.html", results=results, error=error, base=base)


@app.route("/optimize", methods=["GET", "POST"])
def optimize_budget():
    if not authenticated():
        return redirect(url_for("login"))
    result, error = None, None
    with connect_db() as conn:
        trips=conn.execute("SELECT id,source,destination,travelers,days,transportation,budget FROM travel_history WHERE user_id=? ORDER BY id DESC",
                           (session["user_id"],)).fetchall()
    selected_trip_id=None
    if request.method == "POST":
        try:
            raw_trip_id=request.form.get("trip_id", "").strip()
            selected_trip_id=int(raw_trip_id) if raw_trip_id else None
            linked_trip=None
            if selected_trip_id is not None:
                with connect_db() as conn:
                    linked_trip=conn.execute("SELECT * FROM travel_history WHERE id=? AND user_id=?",
                                             (selected_trip_id,session["user_id"])).fetchone()
                if not linked_trip: raise ValueError
            source, destination = (city_name(linked_trip["source"]),city_name(linked_trip["destination"])) if linked_trip else (city_name(request.form.get("source", "")), city_name(request.form.get("destination", "")))
            travelers, days = (int(linked_trip["travelers"]),int(linked_trip["days"])) if linked_trip else (int(request.form.get("travelers", "")), int(request.form.get("days", "")))
            budget = float(linked_trip["budget"] if linked_trip else request.form.get("budget", ""))
            current_transport = linked_trip["transportation"] if linked_trip else request.form.get("transportation", "Train")
            if not source or not destination or source == destination or not 1 <= travelers <= 100 or not 1 <= days <= 365 or not math.isfinite(budget) or budget <= 0 or current_transport not in {"Bus", "Train", "Car"}:
                raise ValueError
            current = estimate_trip(source, destination, travelers, days, (current_transport,))[current_transport]["total_cost"]
            candidates = []
            for candidate_days in range(days, 0, -1):
                estimates = estimate_trip(source, destination, travelers, candidate_days)
                for mode, value in estimates.items():
                    if value["total_cost"] <= budget:
                        changes = (days-candidate_days) + (mode != current_transport)
                        candidates.append((changes, -candidate_days, value["total_cost"], mode, candidate_days))
            if candidates:
                _, _, optimized, mode, optimized_days = min(candidates)
                result = {"current": current, "budget": budget, "optimized": optimized,
                          "saving": max(0, current-optimized), "transport": mode, "days": optimized_days,
                          "changes": [change for change in ([f"Use {mode} instead of {current_transport}"] if mode != current_transport else []) + ([f"Reduce trip duration from {days} to {optimized_days} days"] if optimized_days != days else [])]}
                if current > optimized:
                    with connect_db() as conn:
                        if selected_trip_id is not None:
                            conn.execute("DELETE FROM user_savings WHERE user_id=? AND trip_id=?",
                                         (session["user_id"],selected_trip_id))
                        conn.execute("INSERT INTO user_savings(user_id,trip_id,original_cost,optimized_cost,saved_amount) VALUES(?,?,?,?,?)",
                                     (session["user_id"],selected_trip_id,current,optimized,current-optimized))
                        if selected_trip_id is not None:
                            refresh_actual_savings(conn,session["user_id"],selected_trip_id)
            else:
                result = {"current": current, "budget": budget, "optimized": None,
                          "message": "No modeled transport and duration combination fits this budget. Increase the budget or review the plan."}
        except (ValueError, TypeError, RuntimeError):
            error = "Enter valid trip details and ensure a model is available."
    return render_template("optimize.html", result=result, error=error,trips=trips,selected_trip_id=selected_trip_id)


@app.route("/trips/<int:trip_id>/edit", methods=["POST"])
def edit_trip(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    try:
        source, destination = city_name(request.form.get("source", "")), city_name(request.form.get("destination", ""))
        travelers, days = int(request.form.get("travelers", "")), int(request.form.get("days", ""))
        budget = float(request.form.get("budget", ""))
        transport = request.form.get("transportation", "")
        start_raw=request.form.get("start_date","").strip();end_raw=request.form.get("end_date","").strip()
        start_date=date.fromisoformat(start_raw) if start_raw else None;end_date=date.fromisoformat(end_raw) if end_raw else None
        if bool(start_date)!=bool(end_date) or (start_date and end_date<start_date): raise ValueError
        if start_date: days=(end_date-start_date).days+1
        if not source or not destination or source == destination or len(source)>100 or len(destination)>100 or not 1<=travelers<=100 or not 1<=days<=365 or not math.isfinite(budget) or budget<=0 or transport not in {"Bus","Train","Car"}:
            raise ValueError
        parts = estimate_trip(source,destination,travelers,days,(transport,))[transport]
    except (ValueError, TypeError, RuntimeError):
        return "Enter valid trip details and ensure a model is available.", 400
    with connect_db() as conn:
        row = conn.execute("SELECT id FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"])).fetchone()
        if not row:
            return "Trip not found",404
        conn.execute("""UPDATE travel_history SET source=?,destination=?,travelers=?,days=?,transportation=?,
            transportation_cost=?,accommodation_cost=?,food_cost=?,other_expenses=?,predicted_expense=?,budget=?,start_date=?,end_date=? WHERE id=? AND user_id=?""",
            (source,destination,travelers,days,transport,parts["transport_cost"],parts["accommodation_cost"],parts["food_cost"],parts["other_expenses"],parts["total_cost"],budget,start_date.isoformat() if start_date else None,end_date.isoformat() if end_date else None,trip_id,session["user_id"]))
    return redirect(url_for("trip_detail",trip_id=trip_id))


@app.route("/trips/<int:trip_id>/delete", methods=["POST"])
def delete_trip(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        row = conn.execute("SELECT id FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"])).fetchone()
        if not row:
            return "Trip not found",404
        conn.execute("DELETE FROM expenses WHERE trip_id=? AND user_id=?",(trip_id,session["user_id"]))
        conn.execute("DELETE FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"]))
    return redirect(url_for("history"))


@app.route("/trips/<int:trip_id>/duplicate", methods=["POST"])
def duplicate_trip(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        trip=conn.execute("SELECT * FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"])).fetchone()
        if not trip:
            return "Trip not found",404
        conn.execute("""INSERT INTO travel_history(user_id,source,destination,travelers,days,transportation,
            transportation_cost,accommodation_cost,food_cost,other_expenses,predicted_expense,budget)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (session["user_id"],trip["source"],trip["destination"],trip["travelers"],trip["days"],trip["transportation"],
             trip["transportation_cost"],trip["accommodation_cost"],trip["food_cost"],trip["other_expenses"],trip["predicted_expense"],trip["budget"]))
        new_id=conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    return redirect(url_for("trip_detail",trip_id=new_id))


@app.route("/profile", methods=["GET", "POST"])
def profile():
    if not authenticated():
        return redirect(url_for("login"))
    error = None
    with connect_db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id=?",(session["user_id"],)).fetchone()
    if request.method == "POST":
        name = request.form.get("name","").strip()
        email = request.form.get("email","").strip().lower()
        preferred = request.form.get("preferred_transport","").strip()
        currency = request.form.get("currency","INR").strip().upper()
        raw_budget = request.form.get("default_budget","").strip()
        new_password = request.form.get("new_password","")
        current_password = request.form.get("current_password","")
        theme=request.form.get("theme","light")
        image_file=request.files.get("profile_image")
        new_image=None
        try:
            budget = float(raw_budget) if raw_budget else None
            if not name or len(name)>100 or "@" not in email or len(email)>254 or preferred not in {"", "Bus", "Train", "Car"} or not re.fullmatch(r"[A-Z]{3}",currency) or theme not in {"light","dark"} or (budget is not None and (not math.isfinite(budget) or budget<0)) or (new_password and (len(new_password)<8 or not check_password_hash(user["password"],current_password))):
                raise ValueError
            if image_file and image_file.filename:
                raw=image_file.read(2*1024*1024+1)
                if len(raw)>2*1024*1024: raise ValueError("Profile image exceeds 2 MB.")
                from io import BytesIO
                with Image.open(BytesIO(raw)) as check_image:
                    if check_image.format not in {"PNG","JPEG","WEBP"} or check_image.width*check_image.height>20_000_000:
                        raise ValueError("Use a PNG, JPEG, or WEBP image up to 20 megapixels.")
                    check_image.verify()
                with Image.open(BytesIO(raw)) as source_image:
                    safe_image=ImageOps.exif_transpose(source_image)
                    safe_image=safe_image.convert("RGBA" if "A" in safe_image.getbands() else "RGB")
                    image_dir=BASE_DIR/"static"/"uploads"/"profiles"
                    image_dir.mkdir(parents=True,exist_ok=True)
                    new_image=f"{uuid.uuid4().hex}.png"
                    safe_image.save(image_dir/new_image,format="PNG",optimize=True)
            profile_image=new_image or user["profile_image"]
            with connect_db() as conn:
                if new_password:
                    conn.execute("UPDATE users SET name=?,email=?,preferred_transport=?,default_budget=?,currency=?,theme=?,profile_image=?,password=? WHERE id=?",
                      (name,email,preferred,budget,currency,theme,profile_image,generate_password_hash(new_password),session["user_id"]))
                else:
                    conn.execute("UPDATE users SET name=?,email=?,preferred_transport=?,default_budget=?,currency=?,theme=?,profile_image=? WHERE id=?",
                      (name,email,preferred,budget,currency,theme,profile_image,session["user_id"]))
            session["user"] = name
            session["theme"]=theme
            if new_image and user["profile_image"] and re.fullmatch(r"[a-f0-9]{32}\.png",user["profile_image"]):
                (BASE_DIR/"static"/"uploads"/"profiles"/user["profile_image"]).unlink(missing_ok=True)
            return redirect(url_for("profile"))
        except sqlite3.IntegrityError:
            if new_image: (BASE_DIR/"static"/"uploads"/"profiles"/new_image).unlink(missing_ok=True)
            error="That email is already in use."
        except (ValueError,TypeError,Image.UnidentifiedImageError,Image.DecompressionBombError):
            if new_image: (BASE_DIR/"static"/"uploads"/"profiles"/new_image).unlink(missing_ok=True)
            error = "Check profile fields, image format/size, and password (8 characters minimum)."
    return render_template("profile.html",user=user,error=error)


@app.route("/itinerary", methods=["GET", "POST"])
def itinerary():
    if not authenticated():
        return redirect(url_for("login"))
    plan = None
    if request.method == "POST":
        destination = city_name(request.form.get("destination",""))
        try:
            days = int(request.form.get("days",""))
            budget = float(request.form.get("budget",""))
            if not destination or len(destination)>100 or not 1<=days<=30 or not math.isfinite(budget) or budget<0:
                raise ValueError
            activities = ["Arrival and local orientation", "Visit a major local attraction", "Explore a nearby area or activity", "Flexible time for local experiences", "Shopping and departure"]
            per_day = budget / days if days else 0
            plan = [{"day": i+1,"activity": activities[min((i*len(activities))//days,len(activities)-1)],"budget":round(per_day,2)} for i in range(days)]
        except (ValueError,TypeError):
            plan = []
    return render_template("itinerary.html",plan=plan)


@app.route("/notifications")
def notifications():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        rows=conn.execute("SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 100",(session["user_id"],)).fetchall()
    return render_template("notifications.html",items=rows)


@app.route("/notifications/<int:item_id>/read",methods=["POST"])
def read_notification(item_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        conn.execute("UPDATE notifications SET is_read=1 WHERE id=? AND user_id=?",(item_id,session["user_id"]))
    return redirect(url_for("notifications"))


@app.route("/reports")
def reports():
    if not authenticated():
        return redirect(url_for("login"))
    date_from=request.args.get("date_from","").strip()
    date_to=request.args.get("date_to","").strip()
    error=None
    try:
        if date_from: date.fromisoformat(date_from)
        if date_to: date.fromisoformat(date_to)
        if date_from and date_to and date_from>date_to: raise ValueError
    except ValueError:
        error="Enter a valid date range."
        date_from=date_to=""
    trip_where="user_id=?";trip_params=[session["user_id"]]
    expense_where="user_id=?";expense_params=[session["user_id"]]
    savings_where="user_id=?";savings_params=[session["user_id"]]
    if date_from:
        trip_where+=" AND date(created_at)>=date(?)";trip_params.append(date_from)
        expense_where+=" AND date(expense_date)>=date(?)";expense_params.append(date_from)
        savings_where+=" AND date(created_at)>=date(?)";savings_params.append(date_from)
    if date_to:
        trip_where+=" AND date(created_at)<=date(?)";trip_params.append(date_to)
        expense_where+=" AND date(expense_date)<=date(?)";expense_params.append(date_to)
        savings_where+=" AND date(created_at)<=date(?)";savings_params.append(date_to)
    with connect_db() as conn:
        stats=conn.execute(f"""SELECT COUNT(*),COALESCE(SUM(predicted_expense),0),COALESCE(AVG(predicted_expense),0),
            (SELECT COALESCE(SUM(amount),0) FROM expenses WHERE {expense_where}),
            (SELECT COALESCE(SUM(saved_amount),0) FROM user_savings WHERE {savings_where}),
            (SELECT COALESCE(SUM(actual_savings),0) FROM user_savings WHERE {savings_where}) FROM travel_history WHERE {trip_where}""",
            expense_params+savings_params+savings_params+trip_params).fetchone()
        by_month=conn.execute("""SELECT substr(created_at,1,7) period,SUM(predicted_expense) predicted,0 actual
            FROM travel_history WHERE """+trip_where+""" GROUP BY period
            UNION ALL SELECT substr(expense_date,1,7),0,SUM(amount) FROM expenses WHERE """+expense_where+""" GROUP BY substr(expense_date,1,7)
            ORDER BY period""",trip_params+expense_params).fetchall()
        by_category=conn.execute(f"SELECT category,SUM(amount) amount FROM expenses WHERE {expense_where} GROUP BY category ORDER BY amount DESC",expense_params).fetchall()
    monthly={}
    for row in by_month:
        item=monthly.setdefault(row["period"],{"predicted":0.0,"actual":0.0})
        item["predicted"]+=row["predicted"] or 0
        item["actual"]+=row["actual"] or 0
    return render_template("reports.html",total_trips=stats[0],total_expense=round(stats[1],2),average_expense=round(stats[2],2),actual_expense=round(stats[3],2),savings=round(stats[4],2),actual_savings=round(stats[5],2),months=monthly,categories=by_category,date_from=date_from,date_to=date_to,error=error)


@app.route("/admin/dataset",methods=["GET","POST"])
def admin_dataset():
    denied=require_admin()
    if denied:return denied
    error=summary=None
    if request.method=="POST":
        upload=request.files.get("dataset")
        if not upload or not upload.filename or Path(upload.filename).suffix.lower()!=".csv":
            error="Choose a CSV file."
        else:
            try:
                raw=upload.read(app.config["MAX_CONTENT_LENGTH"]+1)
                if len(raw)>app.config["MAX_CONTENT_LENGTH"]: raise ValueError("File exceeds 5 MB.")
                frame=pd.read_csv(io.BytesIO(raw))
                from train_model import FEATURES,TARGETS
                missing=sorted(set(FEATURES+TARGETS)-set(frame.columns))
                if missing: raise ValueError("Missing required columns: "+", ".join(missing))
                description=frame.describe(include="all").fillna("").round(2)
                summary={"rows":int(len(frame)),"duplicates":int(frame.duplicated().sum()),"missing":int(frame[FEATURES+TARGETS].isna().sum().sum()),"columns":list(frame.columns),"preview":frame.head(8).to_dict(orient="records"),"statistics":description.to_dict(),"stat_names":list(description.index)}
                stored=f"dataset_{uuid.uuid4().hex}.csv"
                folder=BASE_DIR/"dataset"/"uploads";folder.mkdir(parents=True,exist_ok=True)
                (folder/stored).write_bytes(raw)
                with connect_db() as conn:
                    conn.execute("INSERT INTO datasets(stored_name,original_name,row_count,uploaded_by) VALUES(?,?,?,?)",(stored,secure_filename(upload.filename)[:150],len(frame),session["user_id"]))
                    conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                                 (session["user_id"],"dataset_upload",f"{stored}; rows={len(frame)}"))
            except (ValueError,pd.errors.ParserError,UnicodeDecodeError) as exc:
                error=str(exc)
    with connect_db() as conn:
        datasets=conn.execute("SELECT * FROM datasets ORDER BY id DESC").fetchall()
    return render_template("admin_dataset.html",datasets=datasets,error=error,summary=summary)


@app.route("/admin/model/train",methods=["POST"])
def admin_train_model():
    denied=require_admin()
    if denied:return denied
    source=request.form.get("dataset","default")
    if source=="default":
        dataset_path=DATA_PATH
    else:
        with connect_db() as conn:
            row=conn.execute("SELECT stored_name FROM datasets WHERE stored_name=?",(source,)).fetchone()
        if not row: abort(400,"Unknown dataset selection.")
        dataset_path=(BASE_DIR/"dataset"/"uploads"/row["stored_name"]).resolve()
        if dataset_path.parent!=(BASE_DIR/"dataset"/"uploads").resolve(): abort(400)
    try:
        from train_model import train
        with connect_db() as conn:
            number=conn.execute("SELECT COUNT(*) FROM model_versions").fetchone()[0]+1
        version=f"v{number}"
        relative=Path("models")/f"travel_expense_{version}_{uuid.uuid4().hex[:8]}.pkl"
        output=BASE_DIR/relative
        metrics=train(dataset_path,output)
        with connect_db() as conn:
            conn.execute("UPDATE model_versions SET is_active=0")
            conn.execute("INSERT INTO model_versions(version,model_path,training_date,dataset_size,mae,mse,r2_score,rmse,notes,is_active) VALUES(?,?,?,?,?,?,?,?,?,1)",
                (version,relative.as_posix(),metrics["trained_at"],metrics["dataset_size"],metrics["mae"],metrics["mse"],metrics["r2"],metrics["rmse"],f"Dataset: {Path(dataset_path).name}"))
            conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                         (session["user_id"],"model_training",f"{version}; rows={metrics['dataset_size']}"))
        global model,model_metrics
        model,model_metrics=load_model()
        logger.info("Model trained and activated by admin id=%s version=%s",session["user_id"],version)
        return redirect(url_for("admin_models"))
    except Exception as exc:
        logger.exception("Admin model training failed")
        return render_template("admin_dataset.html",datasets=[],error=f"Training failed: {exc}",summary=None),400


@app.route("/admin/models")
def admin_models():
    denied=require_admin()
    if denied:return denied
    with connect_db() as conn:
        rows=conn.execute("SELECT * FROM model_versions ORDER BY id DESC").fetchall()
    return render_template("admin_models.html",models=rows)


@app.route("/admin/models/<int:version_id>/activate",methods=["POST"])
def activate_model(version_id):
    denied=require_admin()
    if denied:return denied
    with connect_db() as conn:
        row=conn.execute("SELECT model_path FROM model_versions WHERE id=?",(version_id,)).fetchone()
    if not row: abort(404)
    path=(BASE_DIR/row["model_path"]).resolve()
    if path.parent!=(BASE_DIR/"models").resolve() or not path.is_file(): abort(400,"Model artifact is missing or invalid.")
    try:
        with path.open("rb") as f: artifact=pickle.load(f)
        if not isinstance(artifact,dict) or "pipeline" not in artifact: abort(400,"Invalid model artifact.")
    except (OSError,pickle.PickleError,EOFError): abort(400,"Invalid model artifact.")
    with connect_db() as conn:
        conn.execute("UPDATE model_versions SET is_active=0")
        conn.execute("UPDATE model_versions SET is_active=1 WHERE id=?",(version_id,))
    global model,model_metrics
    model,model_metrics=load_model()
    return redirect(url_for("admin_models"))


@app.route("/admin/users")
def admin_users():
    denied=require_admin()
    if denied:return denied
    query=request.args.get("q","").strip()[:100]
    with connect_db() as conn:
        users=conn.execute("SELECT id,name,email,role,is_active,created_at FROM users WHERE name LIKE ? OR email LIKE ? ORDER BY id DESC LIMIT 200",(f"%{query}%",f"%{query}%")).fetchall()
    return render_template("admin_users.html",users=users,query=query)


@app.route("/admin/users/<int:user_id>/toggle",methods=["POST"])
def admin_toggle_user(user_id):
    denied=require_admin()
    if denied:return denied
    if user_id==session["user_id"]: abort(400,"You cannot deactivate your own account.")
    with connect_db() as conn:
        user=conn.execute("SELECT role,is_active FROM users WHERE id=?",(user_id,)).fetchone()
        if not user: abort(404)
        if user["role"]=="ADMIN": abort(400,"Admin account status cannot be changed here.")
        conn.execute("UPDATE users SET is_active=? WHERE id=?",(0 if user["is_active"] else 1,user_id))
        conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                     (session["user_id"],"user_status_change",f"user_id={user_id}; active={0 if user['is_active'] else 1}"))
    return redirect(url_for("admin_users"))


@app.route("/admin/logs")
def admin_logs():
    denied=require_admin()
    if denied:return denied
    with connect_db() as conn:
        logs=conn.execute("SELECT l.*,u.email FROM activity_logs l LEFT JOIN users u ON u.id=l.user_id ORDER BY l.id DESC LIMIT 200").fetchall()
    return render_template("admin_logs.html",logs=logs)

@app.route("/compare", methods=["GET","POST"])
def compare_destinations():
    if not authenticated():
        return redirect(url_for("login"))
    rows=[]; error=None
    if request.method=="POST":
        try:
            source=city_name(request.form.get("source",""));days=int(request.form.get("days",""));travelers=int(request.form.get("travelers",""))
            destinations=[city_name(d) for d in request.form.getlist("destinations")]
            if not source or not 1<=days<=365 or not 1<=travelers<=100 or not 2<=len(destinations)<=5 or len(set(destinations))!=len(destinations): raise ValueError
            available=set(city_name(v) for v in pd.read_csv(DATA_PATH)["destination"].dropna().astype(str))
            if any(d not in available for d in destinations): raise ValueError
            for dest in destinations:
                options=estimate_trip(source,dest,travelers,days)
                mode=min(options,key=lambda k:options[k]["total_cost"])
                rows.append({"destination":dest,"transport":mode,**options[mode]})
        except (ValueError,TypeError,RuntimeError,OSError,pd.errors.ParserError):
            error="Select 2â€“5 destinations from the project dataset and enter valid trip details."
    destinations=sorted(set(city_name(v) for v in pd.read_csv(DATA_PATH)["destination"].dropna().astype(str)))
    return render_template("compare.html",rows=rows,error=error,destinations=destinations)

@app.route("/favorites")
def favorites_page():
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        rows=conn.execute("""SELECT f.destination,COUNT(t.id) visits,COALESCE(AVG(t.predicted_expense),0) average_cost,
            MAX(t.created_at) last_visit FROM favorite_destinations f LEFT JOIN travel_history t
            ON t.user_id=f.user_id AND t.destination=f.destination WHERE f.user_id=?
            GROUP BY f.destination ORDER BY visits DESC,f.destination""",(session["user_id"],)).fetchall()
    return render_template("favorites.html",favorites=rows)


@app.route("/favorites/toggle/<int:trip_id>",methods=["POST"])
def toggle_favorite(trip_id):
    if not authenticated():
        return redirect(url_for("login"))
    with connect_db() as conn:
        trip=conn.execute("SELECT destination FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"])).fetchone()
        if not trip:return "Trip not found",404
        found=conn.execute("SELECT 1 FROM favorite_destinations WHERE user_id=? AND destination=?",(session["user_id"],trip["destination"])).fetchone()
        if found: conn.execute("DELETE FROM favorite_destinations WHERE user_id=? AND destination=?",(session["user_id"],trip["destination"]))
        else: conn.execute("INSERT INTO favorite_destinations(user_id,destination) VALUES(?,?)",(session["user_id"],trip["destination"]))
    return redirect(url_for("trip_detail",trip_id=trip_id))

@app.route("/api/predict",methods=["POST"])
def api_predict():
    if not authenticated(): return jsonify({"success":False,"error":"Authentication required."}),401
    payload=request.get_json(silent=True) or {}
    if not isinstance(payload,dict): return jsonify({"success":False,"error":"JSON object required."}),400
    try:
        source,destination=city_name(str(payload.get("source",""))),city_name(str(payload.get("destination","")))
        travelers,days=int(payload.get("travelers")),int(payload.get("days"));budget=float(payload.get("budget"))
        if not source or not destination or source==destination or len(source)>100 or len(destination)>100 or not 1<=travelers<=100 or not 1<=days<=365 or not math.isfinite(budget) or budget<=0: raise ValueError
        predictions=estimate_trip(source,destination,travelers,days)
        affordable={k:v for k,v in predictions.items() if v["total_cost"]<=budget}
        best=min(affordable or predictions,key=lambda key:predictions[key]["total_cost"])
        return jsonify({"success":True,"data":{"predictions":predictions,"recommended_transport":best,"budget":budget,"within_budget":best in affordable}})
    except (ValueError,TypeError,RuntimeError):
        return jsonify({"success":False,"error":"Provide valid trip fields and ensure a model is available."}),400


@app.route("/api/what-if",methods=["POST"])
def api_what_if():
    if not authenticated(): return jsonify({"success":False,"error":"Authentication required."}),401
    payload=request.get_json(silent=True) or {}
    if not isinstance(payload,dict): return jsonify({"success":False,"error":"JSON object required."}),400
    try:
        source,destination=city_name(str(payload.get("source",""))),city_name(str(payload.get("destination","")))
        travelers,days=int(payload.get("travelers")),int(payload.get("days"))
        if not source or not destination or source==destination or not 1<=travelers<=100 or not 1<=days<=365: raise ValueError
        rows=[]
        for mode in ["Bus","Train","Car"]:
            for duration in sorted(set([days,max(1,days-1)])):
                amount=estimate_trip(source,destination,travelers,duration,(mode,))[mode]["total_cost"]
                rows.append({"transportation":mode,"days":duration,"predicted_expense":amount})
        return jsonify({"success":True,"data":{"scenarios":rows}})
    except (ValueError,TypeError,RuntimeError):
        return jsonify({"success":False,"error":"Provide valid trip fields and ensure a model is available."}),400


@app.route("/api/dashboard-stats")
def api_dashboard_stats():
    if not authenticated(): return jsonify({"success":False,"error":"Authentication required."}),401
    with connect_db() as conn:
        stats=conn.execute("""SELECT COUNT(*) trips,COALESCE(SUM(predicted_expense),0) predicted_total,
            COALESCE(AVG(predicted_expense),0) average_prediction,
            (SELECT COALESCE(SUM(amount),0) FROM expenses WHERE user_id=?) actual_total,
            (SELECT COALESCE(SUM(saved_amount),0) FROM user_savings WHERE user_id=?) savings
            FROM travel_history WHERE user_id=?""",(session["user_id"],session["user_id"],session["user_id"])).fetchone()
    return jsonify({"success":True,"data":dict(stats)})
@app.route("/api/csrf")
def api_csrf():
    if not authenticated(): return jsonify({"success":False,"error":"Authentication required."}),401
    if "csrf_token" not in session: session["csrf_token"]=secrets.token_urlsafe(32)
    return jsonify({"success":True,"data":{"csrf_token":session["csrf_token"]}})
@app.route("/reports/export.xlsx")
def export_excel():
    if not authenticated(): return redirect(url_for("login"))
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    with connect_db() as conn:
        trips=conn.execute("""SELECT t.id,t.source,t.destination,t.start_date,t.end_date,t.travelers,t.days,t.transportation,
            t.predicted_expense,t.budget,(SELECT SUM(e.amount) FROM expenses e WHERE e.trip_id=t.id AND e.user_id=t.user_id) actual
            FROM travel_history t WHERE t.user_id=? ORDER BY t.id DESC""",(session["user_id"],)).fetchall()
        expenses=conn.execute("SELECT trip_id,expense_date,category,description,amount FROM expenses WHERE user_id=? ORDER BY expense_date DESC,id DESC",(session["user_id"],)).fetchall()
    book=Workbook();summary=book.active;summary.title="Summary"
    summary.append(["AI Travel Expense Report"]);summary.append(["Metric","Value"])
    summary.append(["Trips",len(trips)]);summary.append(["Predicted total (INR)",sum(float(t["predicted_expense"] or 0) for t in trips)])
    summary.append(["Recorded actual total (INR)",sum(float(e["amount"]) for e in expenses)])
    summary.append(["Optimized potential savings (INR)",sum(float(r[0]) for r in connect_savings(session["user_id"]))])
    trip_sheet=book.create_sheet("Trips");trip_sheet.append(["Trip ID","Source","Destination","Start date","End date","Travelers","Days","Transport","Predicted INR","Budget INR","Actual INR"])
    for row in trips:
        values=[row[k] for k in row.keys()]
        trip_sheet.append([("'"+value if isinstance(value,str) and value[:1] in {"=","+","-","@"} else value) for value in values])
    expense_sheet=book.create_sheet("Actual Expenses");expense_sheet.append(["Trip ID","Date","Category","Description","Amount INR"])
    for row in expenses:
        values=list(row)
        expense_sheet.append([("'"+value if isinstance(value,str) and value[:1] in {"=","+","-","@"} else value) for value in values])
    for ws in book.worksheets:
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="315CE8")
        for column in ws.columns:
            letter=column[0].column_letter;ws.column_dimensions[letter].width=min(36,max(12,max(len(str(c.value or "")) for c in column)+2))
    output=io.BytesIO();book.save(output);output.seek(0)
    return Response(output.getvalue(),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=travel-report.xlsx"})


def connect_savings(user_id):
    with connect_db() as conn:
        return conn.execute("SELECT saved_amount FROM user_savings WHERE user_id=?",(user_id,)).fetchall()

@app.route("/admin/users/<int:user_id>/delete",methods=["POST"])
def admin_delete_user(user_id):
    denied=require_admin()
    if denied:return denied
    if user_id==session["user_id"]:abort(400,"You cannot delete your own account.")
    with connect_db() as conn:
        user=conn.execute("SELECT role,profile_image FROM users WHERE id=?",(user_id,)).fetchone()
        if not user:abort(404)
        if user["role"]=="ADMIN":abort(400,"Admin accounts cannot be deleted from this page.")
        conn.execute("DELETE FROM expenses WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM travel_history WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM wishlist WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM favorite_destinations WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM user_savings WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM notifications WHERE user_id=?",(user_id,))
        conn.execute("UPDATE datasets SET uploaded_by=NULL WHERE uploaded_by=?",(user_id,))
        conn.execute("DELETE FROM activity_logs WHERE user_id=?",(user_id,))
        conn.execute("DELETE FROM users WHERE id=?",(user_id,))
        conn.execute("INSERT INTO activity_logs(user_id,action,details) VALUES(?,?,?)",
                     (session["user_id"],"user_deleted",f"Deleted USER id={user_id}"))
    if user["profile_image"] and re.fullmatch(r"[a-f0-9]{32}\.png",user["profile_image"]):
        (BASE_DIR/"static"/"uploads"/"profiles"/user["profile_image"]).unlink(missing_ok=True)
    return redirect(url_for("admin_users"))

@app.route("/trips/<int:trip_id>/daily-forecast")
def daily_forecast(trip_id):
    if not authenticated():return redirect(url_for("login"))
    with connect_db() as conn:
        trip=conn.execute("SELECT * FROM travel_history WHERE id=? AND user_id=?",(trip_id,session["user_id"])).fetchone()
        if not trip:return "Trip not found",404
        observations=conn.execute("""SELECT t.id,CAST(julianday(e.expense_date)-julianday(t.start_date) AS INTEGER)+1 day_index,
            SUM(e.amount) amount FROM travel_history t JOIN expenses e ON e.trip_id=t.id AND e.user_id=t.user_id
            WHERE t.user_id=? AND t.id<>? AND t.start_date IS NOT NULL AND t.start_date<>''
            AND date(t.end_date)<date('now') AND date(e.expense_date)>=date(t.start_date) AND date(e.expense_date)<=date(t.end_date)
            GROUP BY t.id,day_index ORDER BY t.id,day_index""",(session["user_id"],trip_id)).fetchall()
        actual=conn.execute("""SELECT CAST(julianday(expense_date)-julianday(?) AS INTEGER)+1 day_index,SUM(amount) amount
            FROM expenses WHERE trip_id=? AND user_id=? GROUP BY day_index""",(trip["start_date"] or date.today().isoformat(),trip_id,session["user_id"])).fetchall()
    historical={};trip_ids=set()
    for row in observations:
        if row["day_index"] and row["day_index"]>0:
            historical.setdefault(int(row["day_index"]),[]).append(float(row["amount"]));trip_ids.add(row["id"])
    count=sum(len(values) for values in historical.values())
    enough=count>=3 and len(trip_ids)>=2
    overall=sum(sum(values) for values in historical.values())/count if count else 0
    forecast=[]
    if enough:
        actual_by_day={int(row["day_index"]):float(row["amount"]) for row in actual if row["day_index"] and row["day_index"]>0}
        for day_index in range(1,min(int(trip["days"] or 1),365)+1):
            vals=historical.get(day_index,[])
            forecast.append({"day":day_index,"estimated":round(sum(vals)/len(vals) if vals else overall,2),"actual":actual_by_day.get(day_index)})
    return render_template("daily_forecast.html",trip=trip,forecast=forecast,enough=enough,observations=count,historical_trips=len(trip_ids))

def get_exchange_rate(base,quote):
    base,quote=base.upper(),quote.upper()
    if not re.fullmatch(r"[A-Z]{3}",base) or not re.fullmatch(r"[A-Z]{3}",quote):raise ValueError("Use three-letter currency codes.")
    if base==quote:return 1.0,date.today().isoformat()
    with connect_db() as conn:
        cached=conn.execute("SELECT rate,rate_date FROM exchange_rates WHERE base=? AND quote=? AND fetched_at>=datetime('now','-1 day') ORDER BY fetched_at DESC LIMIT 1",(base,quote)).fetchone()
    if cached:return float(cached["rate"]),cached["rate_date"]
    request_obj=Request(f"https://api.frankfurter.dev/v2/rate/{base}/{quote}",headers={"User-Agent":"TravelExpensePlanner/1.0","Accept":"application/json"})
    with urlopen(request_obj,timeout=5) as response: payload=json.loads(response.read(64*1024))
    rate=float(payload["rate"]);rate_date=str(payload["date"])
    if not math.isfinite(rate) or rate<=0:raise ValueError("The rate service returned invalid data.")
    with connect_db() as conn:
        conn.execute("INSERT INTO exchange_rates(base,quote,rate,rate_date) VALUES(?,?,?,?) ON CONFLICT(base,quote,rate_date) DO UPDATE SET fetched_at=CURRENT_TIMESTAMP,rate=excluded.rate",(base,quote,rate,rate_date))
    return rate,rate_date


@app.route("/calendar")
def travel_calendar():
    if not authenticated():
        return redirect(url_for("login"))
    today=date.today()
    try:
        year=int(request.args.get("year",today.year));month=int(request.args.get("month",today.month))
        if year<2000 or year>2100 or not 1<=month<=12: raise ValueError
    except (ValueError,TypeError):
        year,month=today.year,today.month
    first=date(year,month,1)
    last=date(year,month,calendar_module.monthrange(year,month)[1])
    with connect_db() as conn:
        trips=conn.execute("""SELECT id,source,destination,start_date,end_date,predicted_expense
            FROM travel_history WHERE user_id=? AND start_date IS NOT NULL AND end_date IS NOT NULL
            AND date(start_date)<=date(?) AND date(end_date)>=date(?) ORDER BY start_date,id""",
            (session["user_id"],last.isoformat(),first.isoformat())).fetchall()
    by_day={day:[] for day in range(1,last.day+1)}
    for trip in trips:
        trip_start=max(first,date.fromisoformat(trip["start_date"]))
        trip_end=min(last,date.fromisoformat(trip["end_date"]))
        current=trip_start
        while current<=trip_end:
            by_day[current.day].append(trip)
            current+=timedelta(days=1)
    weeks=[]
    for week in calendar_module.Calendar(firstweekday=0).monthdatescalendar(year,month):
        weeks.append([{"date":day,"in_month":day.month==month,"trips":by_day.get(day.day,[]) if day.month==month else []} for day in week])
    previous=(first-timedelta(days=1))
    next_month=(last+timedelta(days=1))
    return render_template("calendar.html",weeks=weeks,month_name=first.strftime("%B %Y"),
        previous_year=previous.year,previous_month=previous.month,next_year=next_month.year,next_month=next_month.month,
        dated_trip_count=len(trips),today=today)


@app.route("/assistant",methods=["GET","POST"])
def travel_assistant():
    if not authenticated():
        return redirect(url_for("login"))
    question="";answer=None;estimate=None;error=None
    if request.method=="POST":
        question=request.form.get("question","").strip()[:500]
        source=city_name(request.form.get("source",""));destination=city_name(request.form.get("destination",""))
        if source or destination:
            try:
                travelers=int(request.form.get("travelers","1"));days=int(request.form.get("days","1"))
                budget=float(request.form.get("budget","0") or 0);transport=request.form.get("transportation","Train")
                if not source or not destination or source==destination or not 1<=travelers<=100 or not 1<=days<=365 or transport not in {"Bus","Train","Car"} or not math.isfinite(budget) or budget<0:
                    raise ValueError
                estimates=estimate_trip(source,destination,travelers,days)
                chosen=estimates[transport]["total_cost"]
                alternatives=sorted((row["total_cost"],mode) for mode,row in estimates.items())
                best_cost,best_mode=alternatives[0]
                estimate={"source":source,"destination":destination,"days":days,"travelers":travelers,"transport":transport,"cost":chosen,"best_mode":best_mode,"best_cost":best_cost,"budget":budget}
                if budget and best_cost>budget:
                    answer=f"The model estimates {source} to {destination} at ₹{best_cost:,.0f} or more for {days} days. Each supported transport option is above your ₹{budget:,.0f} budget; review a shorter stay or raise the budget."
                elif transport==best_mode:
                    answer=f"The model's lowest estimate is {best_mode} at ₹{best_cost:,.0f}. Your selected {transport} estimate is ₹{chosen:,.0f}."
                else:
                    answer=f"Your selected {transport} estimate is ₹{chosen:,.0f}. The lowest modeled option is {best_mode} at ₹{best_cost:,.0f}."
            except (ValueError,TypeError,RuntimeError):
                error="Enter a different source and destination, 1–100 travelers, 1–365 days, a supported transport, and a valid budget."
        else:
            with connect_db() as conn:
                stats=conn.execute("""SELECT COUNT(*) trips,COALESCE(AVG(predicted_expense),0) average,
                    (SELECT destination FROM travel_history WHERE user_id=? GROUP BY destination ORDER BY COUNT(*) DESC LIMIT 1) favorite,
                    (SELECT destination FROM travel_history WHERE user_id=? AND start_date>=date('now') ORDER BY start_date LIMIT 1) upcoming
                    FROM travel_history WHERE user_id=?""",(session["user_id"],session["user_id"],session["user_id"])).fetchone()
            keywords=question.lower()
            if any(word in keywords for word in ("upcoming","calendar","next trip","scheduled")):
                answer=(f"Your next dated trip is to {stats['upcoming']}. Open the calendar to see its dates." if stats["upcoming"] else "You have no upcoming dated trips. Add departure and return dates while planning a trip, then they will appear on your calendar.")
            elif any(word in keywords for word in ("budget","spend","cost","average","history")):
                answer=(f"You have {stats['trips']} saved trips with an average model estimate of ₹{stats['average']:,.0f}." if stats["trips"] else "There are no saved trips yet. Create a trip plan and I can summarize your travel estimates here.")
                if stats["favorite"]: answer+=f" Your most visited destination is {stats['favorite']}."
            else:
                answer="I can estimate a route with the saved travel model or summarize your trip history. Add a source and destination below for a real estimate, or ask about your budget, spending history, or next trip."
    return render_template("assistant.html",question=question,answer=answer,estimate=estimate,error=error)


@app.route("/currency-converter",methods=["GET","POST"])
def currency_converter():
    if not authenticated():return redirect(url_for("login"))
    result,error=None,None
    if request.method=="POST":
        try:
            amount=float(request.form.get("amount",""));base=request.form.get("base","INR").upper();quote=request.form.get("quote","").upper()
            if not math.isfinite(amount) or amount<0:raise ValueError
            rate,rate_date=get_exchange_rate(base,quote)
            result={"amount":amount,"base":base,"quote":quote,"rate":rate,"converted":round(amount*rate,2),"rate_date":rate_date}
        except (ValueError,TypeError,KeyError,OSError,TimeoutError,json.JSONDecodeError):
            error="Check the amount and currency codes; the current exchange-rate service may be unavailable."
    return render_template("currency.html",result=result,error=error)


@app.route("/api/currency-convert",methods=["POST"])
def api_currency_convert():
    if not authenticated():return jsonify({"success":False,"error":"Authentication required."}),401
    payload=request.get_json(silent=True) or {}
    if not isinstance(payload,dict):return jsonify({"success":False,"error":"JSON object required."}),400
    try:
        amount=float(payload.get("amount"));base=str(payload.get("base","INR"));quote=str(payload.get("quote",""))
        if not math.isfinite(amount) or amount<0:raise ValueError
        rate,rate_date=get_exchange_rate(base,quote)
        return jsonify({"success":True,"data":{"amount":amount,"base":base.upper(),"quote":quote.upper(),"rate":rate,"converted":round(amount*rate,2),"rate_date":rate_date}})
    except (ValueError,TypeError,KeyError,OSError,TimeoutError,json.JSONDecodeError):
        return jsonify({"success":False,"error":"Invalid currency or rate service unavailable."}),503

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1")
